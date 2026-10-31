"""One-time move of the member database and the finance books off Google Sheets.

Reads the two workbooks the team kept in Google Sheets (downloaded as .xlsx —
File → Download → Microsoft Excel) and loads them into the portal:

  members workbook  (MTR_Members_Database)
    Members            -> users (roster fields; sheet is the truth)
    Review - check merges -> users.review_note
    Events             -> events            (matched to existing events by name)
    Participations     -> event_participations

  finance workbook  (MTR_Financial)
    Setup              -> finance_options, finance_seasons, dues tiers/discounts,
                          finance_parties
    Budget             -> budget_lines
    Spendings          -> spendings
    Dues               -> dues_enrollments
    Collection         -> dues_payments
    Funding            -> funding_lines + payables
    Cash               -> cash_transfers

Roster rules (the sheet is the source of truth):
  * each sheet row is matched to an existing member (account) by Gmail, then
    phone, then national ID (only when unique), then name;
  * every MTR ID is re-keyed to the sheet's number (MTR-011, …); portal records
    the sheet doesn't contain keep their row but lose their old MTR ID (the old
    numbering would collide with the new one);
  * a sheet value overwrites the portal's; a blank sheet cell keeps the portal's
    (an existing member keeps their own name);
  * every new member gets an account on the Member level — pre-linked to Google
    when the sheet has their Gmail so they sign straight in, otherwise
    unclaimed (no email, no sign-in) until one is added.

Events rules:
  * matched events get the record columns; their name, status, kind and awards
    are left exactly as the portal has them (the public Hall of Fame reads
    those — awards here would double-count medals the team rows already carry);
  * unmatched events are created; they're never created archived, so nothing
    new appears on the public Hall of Fame until someone archives it on purpose.

Finance: loaded only into empty books. Re-running refuses if the finance tables
already hold data, unless --replace-finance (which wipes and reloads them).

Usage (from backend/):
    python -m scripts.import_member_database MEMBERS.xlsx FINANCE.xlsx           # dry run
    python -m scripts.import_member_database MEMBERS.xlsx FINANCE.xlsx --commit
"""

from __future__ import annotations

import argparse
import re
import sys
import unicodedata
from collections import Counter
from datetime import date, datetime, timezone

import openpyxl
from sqlalchemy import delete, func, select
from sqlalchemy.orm import Session

from app.core.database import SessionLocal
from app.core.security import UNUSABLE_PASSWORD
from app.domains.access.models import AccessLevel
from app.domains.events.models import Event, EventKind, EventParticipation, EventStatus
from app.domains.finance.models import (
    BudgetLine,
    CashTransfer,
    Currency,
    DuesDiscount,
    DuesEnrollment,
    DuesPayment,
    DuesTier,
    EnrollmentStatus,
    FinanceOption,
    FinanceParty,
    FinanceSeason,
    FundingLine,
    FundingStatus,
    OptionList,
    Payable,
    PayableStatus,
    SeasonStatus,
    Spending,
)
from app.domains.members.service import digits, name_key
from app.domains.users.models import User
from app.domains.users.roster import format_mtr_id

# Sheet event IDs whose portal event is named differently.
EVENT_ALIASES = {
    "ROB26": "Robotiáda 2026",
    "FUTECH24": "Fu-Tech Challenge",
    "RCS24": "RCS 2024",
}
# Sheet event "Type" -> portal event kind slug, where the names differ.
KIND_ALIASES = {"general training": "training"}
# The finance books tag budget/spending lines with a competition *series* plus a
# season; this is the event each (series, season) pair means.
SERIES_EVENT = {
    ("MATE ROV", "Season 2026/2027 (S4)"): "ROV27",
    ("Robotiáda", "Season 2026/2027 (S4)"): "ROB26",
}
MEMBER_LEVEL = "Member"


# --- small helpers -------------------------------------------------------------
def s(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    text = str(v).strip()
    return text or None


def num(v) -> float | None:
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return None


def as_int(v) -> int | None:
    n = num(v)
    return int(n) if n is not None else None


def as_date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def as_dt(v) -> datetime | None:
    if isinstance(v, datetime):
        return v.replace(tzinfo=timezone.utc) if v.tzinfo is None else v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day, tzinfo=timezone.utc)
    return None


def fold(text: str | None) -> str:
    """Case/accent-insensitive key: 'Robotiáda' == 'robotiada'."""
    t = unicodedata.normalize("NFKD", text or "")
    return "".join(c for c in t if not unicodedata.combining(c)).strip().lower()


def sheet_rows(ws, header_row: int, first_col: int = 1, last_col: int | None = None) -> list[dict]:
    """Rows below `header_row` as dicts keyed by that row's headers (blank rows skipped)."""
    last_col = last_col or ws.max_column
    headers = [s(ws.cell(header_row, c).value) for c in range(first_col, last_col + 1)]
    out = []
    for r in range(header_row + 1, ws.max_row + 1):
        values = [ws.cell(r, c).value for c in range(first_col, last_col + 1)]
        if all(v in (None, "") for v in values):
            continue
        out.append({h: v for h, v in zip(headers, values) if h})
    return out


class Report:
    def __init__(self) -> None:
        self.counts: Counter = Counter()
        self.notes: list[str] = []

    def add(self, key: str, n: int = 1) -> None:
        self.counts[key] += n

    def note(self, text: str) -> None:
        self.notes.append(text)

    def print(self) -> None:
        for k in sorted(self.counts):
            print(f"  {k:<44}{self.counts[k]:>6}")
        for n in self.notes:
            print(f"  ! {n}")


# --- members -------------------------------------------------------------------
def import_members(db: Session, wb, rep: Report) -> dict[int, User]:
    """Returns {sheet MTR number: member}."""
    ws = wb["Members"]
    rows = []
    for r in sheet_rows(ws, 1):
        number = as_int(r.get("MTR ID"))
        name = s(r.get("Name"))
        if number is None or not name or name.startswith("(reserved"):
            continue
        rows.append((number, r))

    review = {}
    if "Review - check merges" in wb.sheetnames:
        for r in sheet_rows(wb["Review - check merges"], 1):
            n = as_int(r.get("MTR ID"))
            if n is not None:
                why = " — ".join(x for x in (s(r.get("Reason")), s(r.get("Match signals"))) if x)
                review[n] = f"{review[n]}; {why}" if n in review else why

    profiles = list(db.scalars(select(User)))  # every person is a potential match
    users_by_email = {u.email.lower(): u for u in profiles if u.email}
    claimed: set[int] = set()
    matched: dict[int, User] = {}

    def claim(number: int, p: User, how: str) -> None:
        claimed.add(p.id)
        matched[number] = p
        rep.add(f"members matched by {how}")

    # pass 1: Gmail — the strongest key
    by_email = users_by_email
    for number, r in rows:
        e = (s(r.get("Gmail")) or "").lower()
        p = by_email.get(e)
        if e and p is not None and p.id not in claimed:
            claim(number, p, "gmail")

    # pass 2: phone, national ID (unique in the sheet only), name
    nid_counts = Counter(s(r.get("National ID")) for _, r in rows if s(r.get("National ID")))
    for number, r in rows:
        if number in matched:
            continue
        free = [p for p in profiles if p.id not in claimed]
        phone = digits(s(r.get("Phone number")))
        nid = s(r.get("National ID"))
        key = name_key(s(r.get("Name")))
        p = next((p for p in free if phone and digits(p.phone) == phone), None)
        how = "phone"
        if p is None and nid and nid_counts[nid] == 1:
            p, how = next((p for p in free if p.national_id == nid), None), "national id"
        if p is None and key:
            p, how = next((p for p in free if name_key(p.full_name) == key), None), "name"
        if p is not None:
            claim(number, p, how)

    # re-key: clear every MTR ID first so the new numbering never collides
    for p in profiles:
        p.mtr_id = None
    db.flush()
    unmatched_old = [p for p in profiles if p.id not in claimed]
    rep.add("portal members not in the sheet (MTR ID cleared)", len(unmatched_old))

    member_level = db.scalar(select(AccessLevel).where(AccessLevel.name == MEMBER_LEVEL))
    now = datetime.now(timezone.utc)
    for number, r in rows:
        p = matched.get(number)
        gmail = (s(r.get("Gmail")) or "").lower() or None
        if p is None:
            # a new member is a new account: pre-linked to their Gmail so they
            # sign straight in with Google; without one it waits unclaimed
            if gmail in users_by_email:
                rep.add("new members whose Gmail belongs to someone else (email left blank)")
                gmail = None
            p = User(
                email=gmail,
                full_name=s(r.get("Name")),
                hashed_password=UNUSABLE_PASSWORD,
                google_linked_at=now if gmail else None,
                access_level_id=member_level.id if member_level else None,
            )
            db.add(p)
            if gmail:
                users_by_email[gmail] = p
            rep.add("new members (with Gmail)" if gmail else "new members (no Gmail — can't sign in yet)")
        elif gmail and not p.email and gmail not in users_by_email:
            p.email = gmail
            users_by_email[gmail] = p
        # the sheet is the truth for roster fields; an existing account keeps
        # its own name (what the Hall of Fame and website show)
        fields = {
            "phone": s(r.get("Phone number")),
            "national_id": s(r.get("National ID")),
            "birthday": as_date(r.get("Birthday")),
            "location": s(r.get("Governorate")),
            "university": s(r.get("University")),
            "college": s(r.get("College")),
            "major": s(r.get("Major")),
            "uni_id": s(r.get("College ID")),
            "graduating_year": as_int(r.get("Graduating year")),
            "guardian_phone": s(r.get("Guardian's number")),
            "guardian_kinship": s(r.get("Guardian's Kinship")),
            "joined_at": as_dt(r.get("First registration timestamp")),
            "data_confidence": s(r.get("Data confidence")),
            "source_records": as_int(r.get("Source records")),
            "review_note": review.get(number),
        }
        for k, v in fields.items():
            if v is not None:
                setattr(p, k, v)
        p.mtr_id = format_mtr_id(number)
        matched[number] = p
    db.flush()
    rep.add("roster rows in the sheet", len(rows))
    return matched


# --- events --------------------------------------------------------------------
def _kind_for(db: Session, type_name: str | None, rep: Report) -> EventKind | None:
    if not type_name:
        return None
    kinds = list(db.scalars(select(EventKind)))
    want = KIND_ALIASES.get(fold(type_name), fold(type_name))
    for k in kinds:
        if fold(k.slug) == want or fold(k.name) == want:
            return k
    slug = re.sub(r"[^a-z0-9]+", "-", fold(type_name)).strip("-")
    kind = EventKind(slug=slug, name=type_name, sort_order=max((k.sort_order for k in kinds), default=0) + 1)
    db.add(kind)
    db.flush()
    rep.add("event kinds created")
    rep.note(f"created event kind '{type_name}'")
    return kind


def import_events(db: Session, wb, rep: Report) -> dict[str, Event]:
    existing = list(db.scalars(select(Event)))
    by_name = {fold(e.name): e for e in existing}
    by_code = {e.code: e for e in existing if e.code}
    out: dict[str, Event] = {}
    for r in sheet_rows(wb["Events"], 1):
        code = s(r.get("Event ID"))
        name = s(r.get("Event name"))
        if not code or not name:
            continue
        event = by_code.get(code) or by_name.get(fold(EVENT_ALIASES.get(code, name)))
        fields = {
            "code": code,
            "level": s(r.get("Level")),
            "organizer": s(r.get("Organizer")),
            "location": s(r.get("Location")),
            "tracks": s(r.get("Categories / Tracks")),
            "result": s(r.get("Result / Placement")),
            "registered_count": as_int(r.get("# Registered")),
            "participants_count": as_int(r.get("# Participants")),
            "teams_count": as_int(r.get("# Teams")),
            "mentors": s(r.get("Mentors / Coaches")),
            "captains": s(r.get("Team captain(s)")),
            "contact": s(r.get("MTR contact")),
            "budget_total": num(r.get("Budget (total)")),
            "actual_spend": num(r.get("Actual spend")),
            "currency": s(r.get("Currency")),
            "funding_source": s(r.get("Funding source / Sponsor")),
            "registration_form": s(r.get("Registration form")),
            "roster_source": s(r.get("Roster source")),
            "notes": s(r.get("Notes")),
            "prep_date": s(r.get("Competition prep date")),
            "competition_days": s(r.get("Competition day(s)")),
        }
        if event is None:
            kind = _kind_for(db, s(r.get("Type")), rep)
            event = Event(name=name, kind_id=kind.id if kind else None, status=EventStatus.ACTIVE)
            awards = s(r.get("Awards won"))
            event.awards = [awards] if awards else None
            db.add(event)
            rep.add("events created")
            if fold(s(r.get("Status"))) in ("closed", "cancelled"):
                rep.note(f"{code} '{name}' is {s(r.get('Status'))} in the sheet — created active; "
                         "archive it in the portal when it should join the Archive / Hall of Fame")
        else:
            rep.add("events matched")
        for k, v in fields.items():
            if v is not None:
                setattr(event, k, v)
        if as_date(r.get("Start date")) and event.start_date is None:
            event.start_date = as_date(r.get("Start date"))
        if as_date(r.get("End date")) and event.end_date is None:
            event.end_date = as_date(r.get("End date"))
        out[code] = event
    db.flush()
    return out


def import_participations(db: Session, wb, members: dict[int, User],
                          events: dict[str, Event], rep: Report) -> None:
    current = {(p.event_id, p.member_id): p for p in db.scalars(select(EventParticipation))}
    for r in sheet_rows(wb["Participations"], 1):
        number, code = as_int(r.get("MTR ID")), s(r.get("Event ID"))
        member, event = members.get(number), events.get(code)
        if member is None or event is None:
            rep.add("participations skipped (unknown member/event)")
            continue
        p = current.get((event.id, member.id))
        if p is None:
            p = EventParticipation(event_id=event.id, member_id=member.id)
            db.add(p)
            current[(event.id, member.id)] = p
            rep.add("participations created")
        else:
            rep.add("participations updated")
        p.role = s(r.get("Role / Field"))
        p.team = s(r.get("Team"))
        p.mentor = s(r.get("Coach / Mentor"))
        p.rating = s(r.get("Rating"))
        p.registered_at = as_dt(r.get("Registration timestamp"))
    db.flush()


# --- finance -------------------------------------------------------------------
FINANCE_TABLES = [
    DuesPayment, DuesEnrollment, DuesDiscount, DuesTier, CashTransfer, Payable,
    FundingLine, Spending, BudgetLine, FinanceParty, FinanceSeason, FinanceOption,
]


def _column(ws, col: int, start_row: int) -> list[str]:
    out = []
    for r in range(start_row, ws.max_row + 1):
        v = s(ws.cell(r, col).value)
        if v is None:
            break
        out.append(v)
    return out


def import_finance(db: Session, wb, members: dict[int, User],
                   events: dict[str, Event], rep: Report) -> None:
    setup = wb["Setup"]
    options: dict[tuple[str, str], FinanceOption] = {}
    for list_key, col in [
        (OptionList.PAYMENT_METHOD, 1), (OptionList.CATEGORY, 3), (OptionList.SUBCATEGORY, 5),
        (OptionList.PARTY_TYPE, 7), (OptionList.SUBSYSTEM, 19), (OptionList.FUNDING_TYPE, 21),
    ]:
        for i, value in enumerate(_column(setup, col, 5)):
            o = FinanceOption(list_key=list_key, value=value, sort_order=i)
            db.add(o)
            options[(list_key, fold(value))] = o
    db.flush()
    rep.add("finance pick-list values", len(options))

    def opt(list_key: str, value) -> int | None:
        v = s(value)
        if v is None:
            return None
        o = options.get((list_key, fold(v)))
        if o is None:
            o = FinanceOption(list_key=list_key, value=v, sort_order=999)
            db.add(o)
            db.flush()
            options[(list_key, fold(v))] = o
            rep.note(f"added '{v}' to the {list_key} list (used in the books but missing from Setup)")
        return o.id

    settings = {s(setup.cell(r, 23).value): setup.cell(r, 24).value for r in range(5, 30) if s(setup.cell(r, 23).value)}
    usd = num(settings.get("USD → EGP rate")) or 50
    season_start = as_date(settings.get("Season start date"))

    seasons: dict[str, FinanceSeason] = {}
    for i, r in enumerate(range(5, setup.max_row + 1)):
        name = s(setup.cell(r, 13).value)
        if name is None:
            break
        status = fold(s(setup.cell(r, 14).value)) or SeasonStatus.ACTIVE
        start = as_date(setup.cell(r, 15).value)
        season = FinanceSeason(
            name=name,
            status=status if status in {x.value for x in SeasonStatus} else SeasonStatus.ACTIVE,
            start_date=start,
            end_date=as_date(setup.cell(r, 16).value),
            usd_rate=usd,
            notes=s(setup.cell(r, 17).value),
            sort_order=i,
        )
        db.add(season)
        seasons[name] = season
    db.flush()
    # The settings' season start belongs to the first real season.
    if season_start and seasons:
        first = next(iter(seasons.values()))
        first.start_date = first.start_date or season_start
    current = next(iter(seasons.values()), None)
    rep.add("finance seasons", len(seasons))

    if current is not None:
        tiers = {
            "Active": DuesTier(season_id=current.id, name="Active", monthly_fee=num(settings.get("Active base (EGP/mo)")) or 0,
                               months=as_int(settings.get("Active months")) or 0, sort_order=0),
            "Member": DuesTier(season_id=current.id, name="Member", monthly_fee=num(settings.get("Member base (EGP/mo)")) or 0,
                               months=as_int(settings.get("Member months")) or 0, sort_order=1),
            "Mini-robots": DuesTier(season_id=current.id, name="Mini-robots", monthly_fee=0,
                                    months=as_int(settings.get("Mini-robots months")) or 0,
                                    one_time_fee=num(settings.get("Mini-robots one-time (EGP)")) or 0, sort_order=2),
        }
        db.add_all(tiers.values())
        db.add_all([
            DuesDiscount(season_id=current.id, min_years=0, discount=num(settings.get("Discount — new (<1yr)")) or 0, label="New (<1 yr)"),
            DuesDiscount(season_id=current.id, min_years=1, discount=num(settings.get("Discount — 1 year")) or 0, label="1 year"),
            DuesDiscount(season_id=current.id, min_years=2, discount=num(settings.get("Discount — 2 years +")) or 0, label="2 years +"),
        ])
        db.flush()

    # parties ("People & entities"), linked to a portal account by phone
    header_row = next(r for r in range(1, setup.max_row + 1) if s(setup.cell(r, 1).value) == "Name" and s(setup.cell(r, 2).value) == "Type")
    by_phone = {digits(p.phone): p for p in members.values() if digits(p.phone)}
    parties: dict[str, FinanceParty] = {}
    for r in sheet_rows(setup, header_row, 1, 6):
        name = s(r.get("Name"))
        if not name:
            continue
        member = by_phone.get(digits(s(r.get("Phone"))))
        party = FinanceParty(
            name=name, type_id=opt(OptionList.PARTY_TYPE, r.get("Type")), phone=s(r.get("Phone")),
            payment_handle=s(r.get("Payment handle / account")),
            is_active=fold(s(r.get("Active"))) not in ("n", "no", "false"),
            notes=s(r.get("Notes")), user_id=member.id if member else None,
        )
        db.add(party)
        parties[fold(name)] = party
    db.flush()
    rep.add("finance parties", len(parties))

    def party(name) -> int | None:
        n = s(name)
        if n is None:
            return None
        p = parties.get(fold(n))
        if p is None:
            p = FinanceParty(name=n, is_active=True)
            db.add(p)
            db.flush()
            parties[fold(n)] = p
            rep.note(f"added party '{n}' (used in the books but missing from Setup)")
        return p.id

    def season_id(name) -> int | None:
        n = s(name)
        if n is None:
            return current.id if current else None
        return seasons[n].id if n in seasons else (current.id if current else None)

    def event_id(series, season_name) -> int | None:
        code = SERIES_EVENT.get((s(series), s(season_name)))
        if s(series) and code is None:
            rep.note(f"no portal event for '{s(series)}' in '{s(season_name)}' — line kept team-wide")
        return events[code].id if code in events else None

    for r in sheet_rows(wb["Budget"], 4):
        if not s(r.get("Item")):
            continue
        db.add(BudgetLine(
            season_id=season_id(r.get("Season")), event_id=event_id(r.get("Competition"), r.get("Season")),
            subsystem_id=opt(OptionList.SUBSYSTEM, r.get("Subsystem")), item=s(r.get("Item")),
            currency=Currency.USD if s(r.get("Currency")) == "USD" else Currency.EGP,
            unit_price=num(r.get("Unit price (in)")) or 0, required_qty=num(r.get("Required")) or 0,
            spare_qty=num(r.get("Spare")) or 0, in_stock_qty=num(r.get("In-stock")) or 0,
            category_id=opt(OptionList.CATEGORY, r.get("Category")), link_note=s(r.get("Link / Note")),
        ))
        rep.add("budget lines")

    for r in sheet_rows(wb["Spendings"], 4):
        if not s(r.get("Item")):
            continue
        db.add(Spending(
            season_id=season_id(r.get("Season")), event_id=event_id(r.get("Competition"), r.get("Season")),
            spent_on=as_date(r.get("Date")), item=s(r.get("Item")),
            category_id=opt(OptionList.CATEGORY, r.get("Category")),
            subcategory_id=opt(OptionList.SUBCATEGORY, r.get("Sub-category")),
            count=num(r.get("Count")) or 0, unit_cost=num(r.get("Unit cost")) or 0,
            method_id=opt(OptionList.PAYMENT_METHOD, r.get("Method")), paid_by_id=party(r.get("Paid by")),
            receipt=s(r.get("Receipt / Link")), note=s(r.get("Note")),
        ))
        rep.add("spendings")

    by_number = members
    by_name = {name_key(p.full_name): p for p in members.values() if p.full_name}

    def member_of(mtr, name) -> User | None:
        n = as_int(mtr)
        return by_number.get(n) if n is not None and n in by_number else by_name.get(name_key(s(name)))

    if current is not None:
        from app.domains.finance.service import maturity

        enrolled: set[int] = set()
        for r in sheet_rows(wb["Dues"], 17, 1, 9):
            name = s(r.get("Member name"))
            if not name or name.upper() == "TOTAL":
                continue
            m = member_of(r.get("MTR ID"), name)
            if m is None:
                rep.note(f"dues roster: '{name}' is not on the roster — skipped")
                continue
            if m.id in enrolled:
                continue
            tier = tiers.get(s(r.get("Tier")) or "Active", tiers["Active"])
            e = DuesEnrollment(
                season_id=current.id, member_id=m.id, member=m, tier_id=tier.id,
                status=EnrollmentStatus.WITHDRAWN if fold(s(r.get("Status"))) == "withdrawn" else EnrollmentStatus.ACTIVE,
            )
            sheet_maturity = as_int(r.get("Maturity (manual)"))
            if sheet_maturity is not None and maturity(e, current) != sheet_maturity:
                e.maturity_override = sheet_maturity
                rep.add("dues maturity kept from the sheet (differs from join date)")
            db.add(e)
            enrolled.add(m.id)
            rep.add("dues enrolments")

        for r in sheet_rows(wb["Collection"], 4, 1, 7):
            amount, paid_on = num(r.get("Amount (EGP)")), as_date(r.get("Date"))
            if amount is None or paid_on is None:
                continue
            m = member_of(r.get("MTR ID"), r.get("Member"))
            if m is None:
                rep.note(f"collection: '{s(r.get('Member'))}' is not on the roster — skipped")
                continue
            season = max((x for x in seasons.values() if x.start_date and x.start_date <= paid_on),
                         key=lambda x: x.start_date, default=current)
            db.add(DuesPayment(
                season_id=season.id, member_id=m.id, paid_on=paid_on, amount=amount,
                method_id=opt(OptionList.PAYMENT_METHOD, r.get("Method")),
                received_by_id=party(r.get("Received by")), note=s(r.get("Note")),
            ))
            rep.add("dues payments")

    fund = wb["Funding"]
    for r in range(6, fund.max_row + 1):
        label = s(fund.cell(r, 1).value)
        if label == "Money we owe (payables)":
            break
        source = s(fund.cell(r, 2).value)
        if not source or "projected" in source.lower():
            continue  # the projected-dues line is computed live from the dues roster now
        db.add(FundingLine(
            season_id=season_id(label), source=source, type_id=opt(OptionList.FUNDING_TYPE, fund.cell(r, 3).value),
            amount=num(fund.cell(r, 4).value) or 0,
            status=FundingStatus.SECURED if fold(s(fund.cell(r, 5).value)) == "secured" else FundingStatus.EXPECTED,
            note=s(fund.cell(r, 6).value),
        ))
        rep.add("funding lines")
    payables_header = next((r for r in range(1, fund.max_row + 1) if s(fund.cell(r, 1).value) == "Owed To"), None)
    if payables_header:
        for r in sheet_rows(fund, payables_header, 1, 7):
            if not s(r.get("For / description")) and not num(r.get("Amount (EGP)")):
                continue
            db.add(Payable(
                owed_to_id=party(r.get("Owed To")), description=s(r.get("For / description")) or "",
                season_id=season_id(r.get("Season")) if s(r.get("Season")) else None,
                amount=num(r.get("Amount (EGP)")) or 0, due_date=as_date(r.get("Due date")),
                status=PayableStatus.PAID if fold(s(r.get("Status"))) == "paid" else PayableStatus.OWED,
                note=s(r.get("Note")),
            ))
            rep.add("payables")

    cash = wb["Cash"]
    for r in range(6, cash.max_row + 1):
        if s(cash.cell(r, 1).value) == "Cash balances (auto)":
            break
        amount = num(cash.cell(r, 4).value)
        if amount is None:
            continue
        frm = s(cash.cell(r, 2).value)
        db.add(CashTransfer(
            transferred_on=as_date(cash.cell(r, 1).value),
            from_party_id=None if fold(frm) in ("", "opening") else party(frm),
            from_method_id=opt(OptionList.PAYMENT_METHOD, cash.cell(r, 3).value),
            amount=amount, fees=num(cash.cell(r, 5).value) or 0,
            to_party_id=party(cash.cell(r, 6).value),
            to_method_id=opt(OptionList.PAYMENT_METHOD, cash.cell(r, 7).value),
            note=s(cash.cell(r, 8).value),
        ))
        rep.add("cash transfers")
    db.flush()


# --- main ------------------------------------------------------------------------
def run(db: Session, members_path: str, finance_path: str | None, replace_finance: bool) -> Report:
    rep = Report()
    mwb = openpyxl.load_workbook(members_path, data_only=True)
    members = import_members(db, mwb, rep)
    events = import_events(db, mwb, rep)
    import_participations(db, mwb, members, events, rep)
    if finance_path:
        has_books = any(db.scalar(select(func.count()).select_from(t)) for t in FINANCE_TABLES)
        if has_books and not replace_finance:
            rep.note("finance books already hold data — skipped (pass --replace-finance to wipe and reload)")
        else:
            for t in FINANCE_TABLES:
                db.execute(delete(t))
            import_finance(db, openpyxl.load_workbook(finance_path, data_only=True), members, events, rep)
    return rep


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("members_xlsx")
    ap.add_argument("finance_xlsx", nargs="?")
    ap.add_argument("--commit", action="store_true", help="write the changes (default: dry run)")
    ap.add_argument("--replace-finance", action="store_true", help="wipe and reload the finance books")
    args = ap.parse_args()
    db = SessionLocal()
    try:
        rep = run(db, args.members_xlsx, args.finance_xlsx, args.replace_finance)
        print("COMMITTED" if args.commit else "DRY RUN — nothing written (pass --commit)")
        rep.print()
        if args.commit:
            db.commit()
        else:
            db.rollback()
    finally:
        db.close()
    return 0


if __name__ == "__main__":
    sys.exit(main())
