"""Data-entry guardrails on the bulk grid: per-row audit history, the recycle
bin, retired rows staying out of reach, edit conflicts, validation metadata,
column visibility, the guided importer and the season lock."""

from datetime import date

import pytest

from app.domains.finance.models import FinanceOption, FinanceParty, FinanceSeason, Spending


@pytest.fixture()
def season(db_session):
    s = FinanceSeason(name="S5", status="active", start_date=date(2026, 9, 1), usd_rate=50)
    db_session.add(s)
    db_session.flush()
    cash = FinanceOption(list_key="payment_method", value="Cash", sort_order=0)
    db_session.add(cash)
    db_session.add(FinanceParty(name="Osama"))
    db_session.commit()
    return s


def _add_spending(admin, season_id, item="Motor", count="2", unit_cost="100"):
    res = admin.post("/api/bulk/spendings", json={
        "rows": [{"id": "", "season_id": str(season_id), "item": item, "count": count, "unit_cost": unit_cost}],
        "partial": True,
    }).json()
    assert res["ok"], res
    return res["ids"]["0"]


# --- phase 1: audit history + recycle bin --------------------------------------
def test_update_writes_field_level_history(login, org, season):
    admin = login("admin")
    sid = _add_spending(admin, season.id)
    admin.post("/api/bulk/spendings", json={"rows": [{"id": str(sid), "count": "3"}], "partial": True})
    hist = admin.get(f"/api/bulk/spendings/{sid}/history").json()
    assert [h["action"] for h in hist] == ["update", "create"]
    assert hist[0]["detail"]["changes"] == {"count": ["2", "3"]}
    assert hist[0]["actor"]
    assert hist[1]["detail"]["row"]["item"] == "Motor"


def test_hard_delete_goes_to_recycle_bin_and_restores(login, org, season, db_session):
    admin = login("admin")
    sid = _add_spending(admin, season.id, item="Wheel")
    assert admin.post("/api/bulk/spendings", json={"deletes": [sid]}).json()["ok"]
    assert db_session.get(Spending, sid) is None

    bin_ = admin.get("/api/bulk/spendings/trash").json()
    assert len(bin_) == 1 and bin_[0]["detail"]["row"]["item"] == "Wheel"
    res = admin.post(f"/api/bulk/spendings/restore/{bin_[0]['id']}")
    assert res.status_code == 200, res.text
    db_session.expire_all()
    restored = db_session.get(Spending, sid)
    assert restored is not None and restored.item == "Wheel" and restored.count == 2
    assert admin.get("/api/bulk/spendings/trash").json() == []
    # a second restore of the same entry is refused
    assert admin.post(f"/api/bulk/spendings/restore/{bin_[0]['id']}").status_code == 409


def test_restore_refused_when_reference_is_gone(login, org, season, db_session):
    admin = login("admin")
    s2 = FinanceSeason(name="Tmp", status="planning")
    db_session.add(s2)
    db_session.commit()
    sid = _add_spending(admin, s2.id, item="Gone")
    admin.post("/api/bulk/spendings", json={"deletes": [sid]})
    admin.post("/api/bulk/finance_seasons", json={"deletes": [s2.id]})
    entry = next(t for t in admin.get("/api/bulk/spendings/trash").json() if t["entity_id"] == sid)
    res = admin.post(f"/api/bulk/spendings/restore/{entry['id']}")
    assert res.status_code == 409 and "restore" in res.json()["detail"].lower()


def test_soft_deleted_rows_are_out_of_reach_and_restorable(login, org, db_session):
    from app.domains.inventory.models import InventoryItem

    admin = login("admin")
    iid = admin.post("/api/bulk/inventory_items", json={
        "rows": [{"id": "", "name": "Servo", "quantity": "4"}], "partial": True}).json()["ids"]["0"]
    admin.post("/api/bulk/inventory_items", json={"deletes": [iid]})
    # a retired row can't be edited and isn't counted
    res = admin.post("/api/bulk/inventory_items", json={"rows": [{"id": str(iid), "quantity": "9"}], "partial": True}).json()
    assert not res["ok"]
    counts = {t["key"]: t["row_count"] for t in admin.get("/api/bulk/tables").json()}
    assert counts["inventory_items"] == 0
    entry = admin.get("/api/bulk/inventory_items/trash").json()[0]
    assert admin.post(f"/api/bulk/inventory_items/restore/{entry['id']}").status_code == 200
    db_session.expire_all()
    assert db_session.get(InventoryItem, iid).deleted_at is None


def test_history_and_trash_need_table_privilege(login, org, season):
    sid = _add_spending(login("admin"), season.id)
    assert login("ceo").get(f"/api/bulk/spendings/{sid}/history").status_code == 403
    assert login("ceo").get("/api/bulk/spendings/trash").status_code == 403


# --- phase 2: per-cell conflicts ------------------------------------------------
def test_stale_cell_is_a_conflict_but_other_cells_save(login, org, season, db_session):
    admin = login("admin")
    sid = _add_spending(admin, season.id, item="Bolt", count="2")
    # someone else changes count 2 -> 5
    admin.post("/api/bulk/spendings", json={"rows": [{"id": str(sid), "count": "5"}], "partial": True})
    # our editor still thinks count is 2, and also edits the note
    res = admin.post("/api/bulk/spendings", json={"rows": [{
        "id": str(sid), "count": "7", "note": "checked", "_base": {"count": "2", "note": ""},
    }], "partial": True}).json()
    conflict = next(e for e in res["errors"] if e["column"] == "count")
    assert conflict["code"] == "conflict" and conflict["current"] == "5" and conflict["by"]
    db_session.expire_all()
    row = db_session.get(Spending, sid)
    assert row.count == 5 and row.note == "checked"


def test_matching_base_or_same_value_is_not_a_conflict(login, org, season):
    admin = login("admin")
    sid = _add_spending(admin, season.id, count="2")
    ok = admin.post("/api/bulk/spendings", json={"rows": [{"id": str(sid), "count": "3", "_base": {"count": "2"}}],
                                                  "partial": True}).json()
    assert ok["ok"], ok
    # stale base, but the cell already holds what we're writing
    same = admin.post("/api/bulk/spendings", json={"rows": [{"id": str(sid), "count": "3", "_base": {"count": "2"}}],
                                                    "partial": True}).json()
    assert same["ok"], same


# --- phase 3: input rules ---------------------------------------------------------
def test_input_rules_are_enforced_and_described(login, org, season, db_session):
    from app.domains.users.models import MemberProfile

    admin = login("admin")
    cols = {c["name"]: c for c in admin.get("/api/bulk/member_profiles").json()["columns"]}
    assert cols["mtr_id"]["pattern"] and cols["mtr_id"]["pattern_hint"]
    res = admin.post("/api/bulk/spendings", json={"rows": [
        {"id": "", "season_id": str(season.id), "item": "X", "count": "-1"}], "partial": True}).json()
    assert res["errors"][0]["column"] == "count" and "at least 0" in res["errors"][0]["message"]
    res = admin.post("/api/bulk/member_profiles", json={"rows": [
        {"id": "", "full_name": "A", "mtr_id": "123"}], "partial": True}).json()
    assert res["errors"][0]["column"] == "mtr_id"
    # a legacy value that breaks a rule can be sent back unchanged
    m = MemberProfile(full_name="Legacy", phone="phone?")
    db_session.add(m)
    db_session.commit()
    ok = admin.post("/api/bulk/member_profiles", json={"rows": [
        {"id": str(m.id), "phone": "phone?", "university": "AUC"}], "partial": True}).json()
    assert ok["ok"], ok


# --- phase 4: column-level visibility ---------------------------------------------
def _grant(db, login_name, org, *privs):
    import json

    from app.domains.access.models import AccessLevel

    level = db.get(AccessLevel, org[login_name].access_level_id)
    level.privileges = json.dumps(sorted(set(json.loads(level.privileges)) | set(privs)))
    db.commit()


def test_member_contact_columns_need_members_edit(login, org, db_session):
    from app.domains.users.models import MemberProfile

    m = MemberProfile(full_name="Private Pat", mtr_id="MTR-500", phone="01012345678", national_id="12345678901234")
    db_session.add(m)
    db_session.commit()
    _grant(db_session, "ceo", org, "dashboard.view", "members.view")
    ceo = login("ceo")
    rec = ceo.get(f"/api/explore/member_profiles/{m.id}").json()
    names = {f["name"] for f in rec["fields"]}
    assert "mtr_id" in names and "phone" not in names and "national_id" not in names
    # can't find someone by a number they aren't allowed to see
    assert ceo.get("/api/explore/member_profiles", params={"q": "0101234"}).json()["total"] == 0

    admin_rec = login("admin").get(f"/api/explore/member_profiles/{m.id}").json()
    assert "phone" in {f["name"] for f in admin_rec["fields"]}


def test_narrowed_spec_drops_columns_and_values():
    from app.domains.bulk.registry import TABLES

    spec = TABLES["member_profiles"].for_privileges(set())
    assert "phone" not in {c.name for c in spec.columns}
    assert spec.full is TABLES["member_profiles"]
    assert TABLES["member_profiles"].for_privileges({"members.edit"}) is TABLES["member_profiles"]


def test_members_directory_redacts_contact_details(login, org, db_session):
    from app.domains.users.models import MemberProfile

    m = MemberProfile(full_name="Quiet Q", mtr_id="MTR-501", phone="01099998888", guardian_phone="01011112222")
    db_session.add(m)
    db_session.commit()
    _grant(db_session, "ceo", org, "members.view")
    ceo = login("ceo")
    assert next(r for r in ceo.get("/api/members").json() if r["id"] == m.id)["phone"] is None
    detail = ceo.get(f"/api/members/{m.id}").json()
    assert detail["phone"] is None and detail["guardian_phone"] is None and detail["mtr_id"] == "MTR-501"
    assert login("admin").get(f"/api/members/{m.id}").json()["phone"] == "01099998888"


# --- phase 5: guided importer -------------------------------------------------------
def _book(content):
    import io

    import openpyxl

    return openpyxl.load_workbook(io.BytesIO(content))


def _save(wb) -> bytes:
    import io

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _upload(client, table, content, query="", name="f.xlsx", data=None):
    return client.post(f"/api/bulk/{table}/upload{query}", files={"file": (name, content, "application/octet-stream")},
                       data=data or {})


def test_export_is_scoped_typed_and_has_dropdowns(login, org, season, db_session):
    admin = login("admin")
    other = FinanceSeason(name="Old", status="closed")
    db_session.add(other)
    db_session.commit()
    _add_spending(admin, season.id, item="In scope")
    _add_spending(admin, other.id, item="Out of scope")
    wb = _book(admin.get(f"/api/bulk/spendings/export.xlsx?season_id={season.id}").content)
    ws = wb["spendings"]
    items = [r[4] for r in ws.iter_rows(min_row=2, values_only=True)]
    assert items == ["In scope"]
    header = [c.value for c in ws[1]]
    assert header[0] == "ID" and header[-1] == "Status"
    row = next(ws.iter_rows(min_row=2, values_only=True))
    assert row[header.index("Season")] == "S5"  # a label, not an id
    assert isinstance(row[header.index("Count")], (int, float))  # a real number: formulas work
    assert ws.data_validations.dataValidation  # dropdowns
    assert wb["_meta"].sheet_state == "hidden"


def test_untouched_cells_are_not_written_and_stale_cells_conflict(login, org, season, db_session):
    admin = login("admin")
    sid = _add_spending(admin, season.id, item="Gear", count="2")
    content = admin.get(f"/api/bulk/spendings/export.xlsx?season_id={season.id}").content
    # meanwhile someone edits count in the portal
    admin.post("/api/bulk/spendings", json={"rows": [{"id": str(sid), "count": "4"}], "partial": True})

    # the file only changes the note: applies cleanly and keeps count=4
    wb = _book(content)
    ws = wb["spendings"]
    hdr = [c.value for c in ws[1]]
    ws.cell(row=2, column=hdr.index("Note") + 1, value="from excel")
    res = _upload(admin, "spendings", _save(wb), "?apply=true").json()
    assert res["applied"], res
    db_session.expire_all()
    row = db_session.get(Spending, sid)
    assert row.count == 4 and row.note == "from excel"

    # a file that changes the stale count is a conflict, nothing written
    wb = _book(content)
    wb["spendings"].cell(row=2, column=hdr.index("Count") + 1, value=9)
    prev = _upload(admin, "spendings", _save(wb)).json()
    assert not prev["ok"] and prev["errors"][0]["code"] == "conflict"


def test_foreign_file_headers_labels_and_mapping(login, org, season, db_session):
    import openpyxl

    admin = login("admin")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Item", "season", "Unit Cost", "Qty bought", "Paid with"])
    ws.append(["Battery", "S5", 300, 2, "cash"])
    content = _save(wb)
    prev = _upload(admin, "spendings", content).json()
    assert prev["unmatched"] == ["Qty bought", "Paid with"]
    import json

    prev = _upload(admin, "spendings", content, data={"mapping": json.dumps({"Qty bought": "count", "Paid with": "method_id"})}).json()
    assert prev["ok"] and prev["new"] == 1, prev
    res = _upload(admin, "spendings", content, "?apply=true",
                  data={"mapping": json.dumps({"Qty bought": "count", "Paid with": "method_id"})}).json()
    assert res["applied"]
    sp = db_session.query(Spending).filter_by(item="Battery").one()
    assert sp.count == 2 and sp.method_id is not None and sp.season_id == season.id


def test_scoped_delete_missing_never_touches_other_scopes(login, org, season, db_session):
    admin = login("admin")
    other = FinanceSeason(name="Other", status="active")
    db_session.add(other)
    db_session.commit()
    keep_other = _add_spending(admin, other.id, item="Other season")
    _add_spending(admin, season.id, item="Delete me")
    wb = _book(admin.get(f"/api/bulk/spendings/export.xlsx?season_id={season.id}").content)
    wb["spendings"].delete_rows(2)
    prev = _upload(admin, "spendings", _save(wb), "?delete_missing=true").json()
    assert prev["will_delete"] == 1 and keep_other not in prev["missing_ids"]


def test_annotated_workbook_marks_bad_rows(login, org, season):
    import openpyxl

    admin = login("admin")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Season", "Item", "Count"])
    ws.append(["S5", "Fine", 1])
    ws.append(["S5", "Bad", -3])
    r = _upload(admin, "spendings", _save(wb), "?annotate=true")
    assert r.status_code == 200
    out = _book(r.content).active
    status = [row[-1] for row in out.iter_rows(min_row=2, values_only=True)]
    assert status[0] == "✓" and "at least 0" in status[1]


def test_upload_row_cap(login, org, season, monkeypatch):
    import openpyxl

    from app.domains.bulk import service

    monkeypatch.setattr(service, "UPLOAD_MAX_ROWS", 2)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Season", "Item"])
    for i in range(3):
        ws.append(["S5", f"x{i}"])
    r = _upload(login("admin"), "spendings", _save(wb))
    assert r.status_code == 400 and "more than 2 rows" in r.json()["detail"]
