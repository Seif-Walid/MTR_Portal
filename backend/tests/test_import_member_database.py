"""The one-time Sheets -> portal import: roster matching + MTR ID re-keying,
roster-only members, events matched without touching Hall-of-Fame fields,
participations, and idempotency."""

from datetime import datetime

import openpyxl

from app.domains.events.models import Event, EventParticipation, EventStatus
from app.domains.users.models import MemberProfile, User
from scripts import import_member_database as imp

MEMBER_HEADERS = ["MTR ID", "Name", "Phone number", "National ID", "Birthday", "Governorate",
                  "University", "College", "Major", "College ID", "Graduating year", "Gmail",
                  "Guardian's number", "Guardian's Kinship", "First registration timestamp",
                  "Events attended", "Data confidence", "Source records"]


def _workbook(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Members"
    ws.append(MEMBER_HEADERS)
    ws.append([2, "(reserved for special member)"])
    ws.append([11, "Student Sheet Name", "01200000001", None, None, "Alexandria", "ANU", None, "CCE",
               None, "2028", "stud@t.local", None, None, datetime(2024, 7, 1), 1, "high", 3])
    ws.append([12, "No Gmail Person", "01200000002", None, None, None, "ANU", None, "MRE",
               None, "2029", None, None, None, datetime(2025, 7, 1), 1, "single", 1])
    ws.append([13, "Brand New", "01200000003", None, None, None, "AASTMT", None, None,
               None, None, "brandnew@gmail.com", None, None, datetime(2026, 7, 1), 0, "single", 1])
    ev = wb.create_sheet("Events")
    ev.append(["Event ID", "Event name", "Type", "Level", "Status", "Result / Placement", "Awards won", "# Participants"])
    ev.append(["MRC25", "MRC 2025", "Competition", "National", "Closed", "Sumo: 1st", "🥇 1st — Sumo", 12])
    ev.append(["IRCCHN25", "IRC China 2025", "Competition", "International", "Closed", "Sumo: 5th", None, None])
    ev.append(["GT-BOOT26", "Bootcamp Summer 2026", "Bootcamp", None, None, None, None, 27])
    pa = wb.create_sheet("Participations")
    pa.append(["MTR ID", "Name", "Event ID", "Event", "Event type", "Role / Field", "Team", "Coach / Mentor", "Rating", "Registration timestamp"])
    pa.append([11, "x", "MRC25", "MRC 2025", "Competition", "Software", "Line 1", "Osama", "95.5", None])
    pa.append([12, "x", "GT-BOOT26", "Bootcamp", "Bootcamp", None, None, None, None, None])
    pa.append([99, "unknown", "MRC25", "MRC 2025", "Competition", None, None, None, None, None])
    path = tmp_path / "members.xlsx"
    wb.save(path)
    return str(path)


def test_import_roster_events_participations(db_session, org, tmp_path):
    db = db_session
    from app.domains.events.service import ensure_preset_kinds  # noqa: F401 — org fixture seeded kinds
    student = org["student"]
    db.add(MemberProfile(user_id=student.id, mtr_id="MTR-099", location="Cairo"))
    archived = Event(name="MRC 2025", status=EventStatus.ARCHIVED, awards=["🥇 curated"])
    db.add(archived)
    db.commit()

    path = _workbook(tmp_path)
    rep = imp.run(db, path, None, False)
    db.commit()

    by_mtr = {p.mtr_id: p for p in db.query(MemberProfile)}
    # matched by Gmail, re-keyed to the sheet number, sheet values win, blanks keep
    assert by_mtr["MTR-011"].user_id == student.id
    assert by_mtr["MTR-011"].location == "Alexandria" and by_mtr["MTR-011"].major == "CCE"
    assert by_mtr["MTR-011"].graduating_year == 2028
    assert "MTR-099" not in by_mtr
    # roster-only member: no Gmail, no account
    assert by_mtr["MTR-012"].user_id is None and by_mtr["MTR-012"].full_name == "No Gmail Person"
    # new member with a Gmail gets a Google-linked account
    new_user = db.query(User).filter_by(email="brandnew@gmail.com").one()
    assert by_mtr["MTR-013"].user_id == new_user.id and new_user.google_linked_at is not None

    # matched event keeps its curated name/status/awards, gains the record fields
    db.refresh(archived)
    assert archived.code == "MRC25" and archived.result == "Sumo: 1st" and archived.participants_count == 12
    assert archived.status == EventStatus.ARCHIVED and archived.awards == ["🥇 curated"]
    # unmatched events are created, never archived (nothing new goes public)
    china = db.query(Event).filter_by(code="IRCCHN25").one()
    assert china.status == EventStatus.ACTIVE
    boot = db.query(Event).filter_by(code="GT-BOOT26").one()
    assert boot.kind is not None and boot.kind.name == "Bootcamp"

    parts = db.query(EventParticipation).all()
    assert len(parts) == 2 and rep.counts["participations skipped (unknown member/event)"] == 1
    assert next(p for p in parts if p.event_id == archived.id).team == "Line 1"

    # idempotent: a second run creates nothing new
    rep2 = imp.run(db, path, None, False)
    db.commit()
    assert db.query(EventParticipation).count() == 2
    assert db.query(MemberProfile).filter(MemberProfile.mtr_id.is_not(None)).count() == 3
    assert rep2.counts["events created"] == 0 and rep2.counts["new members (portal account created)"] == 0
